from pathlib import Path
from openpyxl import Workbook, load_workbook
from config.settings import settings
from classification.rule_classifier import PALABRAS_CLAVE

ENCABEZADOS = ["instancia", *PALABRAS_CLAVE]

def registrar_instancia(instancia: int, tipos: dict[str, bool]) -> Path:
    """
    Agrega (o actualiza, si ya existe) la fila de la instancia en el Excel de resultados:
    instancia | SUNAFIL | CUL | BOLETA  con valores SI / NO.
    """
    ruta = Path(settings.RESULTADOS_XLSX)
    ruta.parent.mkdir(parents=True, exist_ok=True)

    if ruta.exists():
        wb = load_workbook(ruta)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Validacion"
        ws.append(ENCABEZADOS)

    fila = [instancia, *("SI" if tipos[t] else "NO" for t in PALABRAS_CLAVE)]

    for row in ws.iter_rows(min_row=2):
        if row[0].value == instancia:
            for celda, valor in zip(row, fila):
                celda.value = valor
            break
    else:
        ws.append(fila)

    try:
        wb.save(ruta)
    except PermissionError:
        raise PermissionError(f"No se pudo guardar {ruta}: ciérralo en Excel y vuelve a ejecutar") from None
    return ruta
