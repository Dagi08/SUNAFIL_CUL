import io
import logging
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document as DocxDocument
from PIL import Image, ImageOps, ImageSequence

from config.settings import settings
from extraction.base_extractor import BaseExtractor
from models.document import Pagina

logger = logging.getLogger(__name__)

EXTENSIONES_IMAGEN = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
# Tesseract falla con imágenes chicas (fotos de celular pegadas en un PDF): se amplían hasta este ancho
ANCHO_MINIMO_OCR = 1800

class OcrExtractor(BaseExtractor):
    """
    Extrae texto por página:
      - PDF: texto nativo con PyMuPDF; si la página es una captura/escaneo, OCR sobre la página renderizada.
      - Imágenes: OCR directo (cada frame de un TIFF es una página).
      - DOCX: texto del documento como página 1 y cada imagen incrustada como una página más (OCR).
    """

    def __init__(self):
        if settings.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
        self.lang = self._resolver_idioma(settings.OCR_LANG)

    def _resolver_idioma(self, lang: str) -> str:
        disponibles = set(pytesseract.get_languages(config=""))
        pedidos = lang.split("+")
        validos = [l for l in pedidos if l in disponibles]
        faltantes = [l for l in pedidos if l not in disponibles]
        if faltantes:
            logger.warning(
                f"Idioma(s) OCR no instalados en Tesseract: {faltantes}. "
                f"Descarga <idioma>.traineddata en la carpeta tessdata. Disponibles: {sorted(disponibles)}"
            )
        return "+".join(validos) or "eng"

    def extraer(self, ruta: Path) -> list[Pagina]:
        extension = ruta.suffix.lower()
        if extension == ".pdf":
            return self._extraer_pdf(ruta)
        if extension in EXTENSIONES_IMAGEN:
            return self._extraer_imagen(ruta)
        if extension == ".docx":
            return self._extraer_docx(ruta)
        raise ValueError(f"Extensión no soportada: {extension}")

    def _ocr(self, imagen: Image.Image) -> str:
        imagen = ImageOps.exif_transpose(imagen).convert("L")  # escala de grises mejora el OCR
        if imagen.width < ANCHO_MINIMO_OCR:
            factor = ANCHO_MINIMO_OCR / imagen.width
            imagen = imagen.resize((ANCHO_MINIMO_OCR, int(imagen.height * factor)), Image.LANCZOS)
        return pytesseract.image_to_string(imagen, lang=self.lang).strip()

    def _extraer_pdf(self, ruta: Path) -> list[Pagina]:
        paginas = []
        with pymupdf.open(ruta) as pdf:
            for i, page in enumerate(pdf, start=1):
                texto = page.get_text().strip()
                if len(texto) >= settings.MIN_CHARS_TEXTO_NATIVO:
                    paginas.append(Pagina(i, "texto_nativo", texto))
                    continue
                pix = page.get_pixmap(dpi=settings.PDF_DPI)
                texto = self._ocr(Image.open(io.BytesIO(pix.tobytes("png"))))
                if len(texto) < settings.MIN_CHARS_OCR_PAGINA:
                    # La página completa dio poco texto: se intenta con cada foto incrustada por separado
                    texto = max(texto, self._ocr_imagenes_incrustadas(pdf, page), key=len)
                paginas.append(Pagina(i, "ocr", texto))
        return paginas

    def _ocr_imagenes_incrustadas(self, pdf: pymupdf.Document, page: pymupdf.Page) -> str:
        textos = []
        for xref, *_ in page.get_images(full=True):
            try:
                datos = pdf.extract_image(xref)["image"]
                textos.append(self._ocr(Image.open(io.BytesIO(datos))))
            except Exception as e:
                logger.warning(f"No se pudo procesar imagen incrustada (xref {xref}): {e}")
        return "\n\n".join(t for t in textos if t)

    def _extraer_imagen(self, ruta: Path) -> list[Pagina]:
        with Image.open(ruta) as imagen:
            return [
                Pagina(i, "ocr", self._ocr(frame.copy()))
                for i, frame in enumerate(ImageSequence.Iterator(imagen), start=1)
            ]

    def _extraer_docx(self, ruta: Path) -> list[Pagina]:
        doc = DocxDocument(ruta)
        lineas = [p.text for p in doc.paragraphs if p.text.strip()]
        for tabla in doc.tables:
            for fila in tabla.rows:
                lineas.append(" | ".join(celda.text.strip() for celda in fila.cells))

        paginas = []
        if lineas:
            paginas.append(Pagina(1, "texto_nativo", "\n".join(lineas)))

        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue
            try:
                imagen = Image.open(io.BytesIO(rel.target_part.blob))
                paginas.append(Pagina(len(paginas) + 1, "ocr", self._ocr(imagen)))
            except Exception as e:
                logger.warning(f"No se pudo procesar imagen incrustada en {ruta.name}: {e}")
        return paginas
